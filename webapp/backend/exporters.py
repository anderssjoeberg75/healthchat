"""Conversation export (TXT / PDF / DOCX).

Same output as the desktop ``ExportReportDialog``, but written to an in-memory
buffer and streamed to the browser as a download instead of being saved through
a file dialog.
"""

import io
import logging
from typing import Any, Dict, List

logger = logging.getLogger(__name__)


def _visible(messages: List[Dict[str, Any]], include_system: bool) -> List[Dict[str, Any]]:
    if include_system:
        return messages
    return [m for m in messages if m.get("type") != "system"]


def export_txt(messages: List[Dict[str, Any]], include_timestamp: bool, include_system: bool) -> bytes:
    buffer = io.StringIO()
    buffer.write("GARMIN CHAT CONVERSATION REPORT\n")
    buffer.write("=" * 60 + "\n\n")

    for message in _visible(messages, include_system):
        sender = message.get("sender", "Unknown")
        if include_timestamp:
            buffer.write(f"[{message.get('timestamp', '')}] {sender}:\n")
        else:
            buffer.write(f"{sender}:\n")
        buffer.write(f"{message.get('message', '')}\n\n")
        buffer.write("-" * 60 + "\n\n")

    return buffer.getvalue().encode("utf-8")


def export_pdf(messages: List[Dict[str, Any]], include_timestamp: bool, include_system: bool) -> bytes:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)
    styles = getSampleStyleSheet()
    story = [Paragraph("HealthChat Conversation Report", styles["Title"]), Spacer(1, 0.3 * inch)]

    for message in _visible(messages, include_system):
        sender = message.get("sender", "Unknown")
        header = (
            f"<b>[{message.get('timestamp', '')}] {sender}:</b>" if include_timestamp else f"<b>{sender}:</b>"
        )
        story.append(Paragraph(header, styles["Normal"]))
        # Escape markup so a stray '<' in an answer cannot break the PDF build.
        text = (
            str(message.get("message", ""))
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
            .replace("\n", "<br/>")
        )
        story.append(Paragraph(text, styles["Normal"]))
        story.append(Spacer(1, 0.2 * inch))

    doc.build(story)
    return buffer.getvalue()


def export_docx(messages: List[Dict[str, Any]], include_timestamp: bool, include_system: bool) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    document = Document()
    title = document.add_heading("HealthChat Conversation Report", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph()

    for message in _visible(messages, include_system):
        sender = message.get("sender", "Unknown")
        paragraph = document.add_paragraph()
        if include_timestamp:
            paragraph.add_run(f"[{message.get('timestamp', '')}] {sender}:").bold = True
        else:
            paragraph.add_run(f"{sender}:").bold = True
        document.add_paragraph(str(message.get("message", "")))
        document.add_paragraph("_" * 60)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


EXPORTERS = {"txt": export_txt, "pdf": export_pdf, "docx": export_docx}
MEDIA_TYPES = {
    "txt": "text/plain; charset=utf-8",
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
